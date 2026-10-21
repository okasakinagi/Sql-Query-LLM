# -*- coding: utf-8 -*-
"""导出为 Word 文档（服务端生成，python-docx）。

图表由前端 ECharts 导出为 PNG（base64），随请求一起提交后嵌入文档，
避免服务端渲染中文图表带来的字体依赖问题。
"""

from __future__ import annotations

import base64
import datetime as dt
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

_CHART_MIME = re.compile(r"^data:image/\w+;base64,", re.I)


def _add_table(doc: Document, columns: list[str], rows: list[list]) -> None:
    if not columns:
        return
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, name in enumerate(columns):
        hdr[i].text = str(name)
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)

    # 小字号，避免宽表溢出页面
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)


def build_docx(payload: dict) -> bytes:
    """根据一次问答结果生成 Word 文档，返回字节流。"""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    # ---------------- 标题 ----------------
    title = doc.add_heading("智能数据库查询报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"生成时间：{dt.datetime.now():%Y-%m-%d %H:%M:%S}    "
        f"数据源：erp_demo（中小企业进销存与销售管理）"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # ---------------- 问题 ----------------
    doc.add_heading("一、查询问题", level=1)
    doc.add_paragraph(payload.get("question", ""))

    # ---------------- 答案 ----------------
    doc.add_heading("二、分析结论", level=1)
    doc.add_paragraph(payload.get("answer", "（无）"))

    # ---------------- 图表 ----------------
    chart_img = payload.get("chartImage") or ""
    if chart_img:
        doc.add_heading("三、统计图表", level=1)
        try:
            raw = _CHART_MIME.sub("", chart_img)
            image = io.BytesIO(base64.b64decode(raw))
            doc.add_picture(image, width=Inches(5.9))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except (ValueError, TypeError):
            doc.add_paragraph("（图表数据格式异常，未能嵌入）")

    # ---------------- SQL ----------------
    doc.add_heading("四、生成的 SQL 语句", level=1)
    sql_par = doc.add_paragraph()
    sql_run = sql_par.add_run(payload.get("sql", ""))
    sql_run.font.name = "Consolas"
    sql_run.font.size = Pt(9)
    sql_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    # ---------------- 明细数据 ----------------
    doc.add_heading("五、查询结果明细", level=1)
    columns = payload.get("columns") or []
    rows = payload.get("rows") or []
    if columns:
        info = doc.add_paragraph(
            f"共 {payload.get('row_count', len(rows))} 行"
            + ("（导出仅包含前 200 行）" if len(rows) > 200 else "")
        )
        info.runs[0].font.size = Pt(9)
        _add_table(doc, columns, rows[:200])
    else:
        doc.add_paragraph("（无结果数据）")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
